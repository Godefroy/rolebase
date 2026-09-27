#!/usr/bin/env python3
"""Generate the PowerPoint org chart templates offered on the "org chart in PowerPoint" post.

Lives beside the page it serves: `en.mdx` and `fr.mdx` in this folder,
published on /en/blog/org-chart-powerpoint and /fr/blog/org-chart-powerpoint.

Files produced, under `public/downloads/` and served on `/downloads/<file>`
    org-chart-template.pptx   EN, linked from en.mdx
    modele-organigramme.pptx  FR, linked from fr.mdx

Each deck holds six 16:9 slides:
    1. the role list to fill in before drawing anything
    2. a filled three-level chart, whose branch boxes link to slides 3 to 5
    3-5. one slide per branch, with a link back to the overview
    6. a blank chart to type over

The boxes are plain shapes joined by elbow connectors rather than SmartArt,
so the reader moves a box and its lines follow.

These binaries are build artefacts committed to the repo. Edit the definitions
below and re-run this script rather than editing the files themselves.

    pip install python-pptx
    python3 src/content/blog/org-chart-powerpoint/generate-org-chart-template.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

OUT = Path(__file__).resolve().parents[4] / "public" / "downloads"

PURPLE = RGBColor(0x98, 0x70, 0xF0)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
LINE = RGBColor(0x9A, 0x9A, 0x9A)
BOX_FILL = RGBColor(0xF3, 0xED, 0xFE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Cm(33.867)
SLIDE_H = Cm(19.05)
MARGIN = Cm(1.5)
BOX_W = Cm(4.6)
BOX_H = Cm(1.9)

# connection points of a rectangle, as PowerPoint numbers them
TOP, BOTTOM = 0, 2


def text(frame, value, size, bold=False, colour=None, align=PP_ALIGN.LEFT, italic=False):
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if colour is not None:
        run.font.color.rgb = colour
    return p


def title(slide, heading, hint):
    box = slide.shapes.add_textbox(MARGIN, Cm(0.8), SLIDE_W - 2 * MARGIN, Cm(1.4))
    text(box.text_frame, heading, 26, bold=True, colour=PURPLE)
    box = slide.shapes.add_textbox(MARGIN, Cm(2.3), SLIDE_W - 2 * MARGIN, Cm(1.6))
    text(box.text_frame, hint, 13, colour=GREY, italic=True)


def footer(slide, value):
    box = slide.shapes.add_textbox(MARGIN, SLIDE_H - Cm(1.3), SLIDE_W - 2 * MARGIN, Cm(0.8))
    text(box.text_frame, value, 10, colour=GREY, align=PP_ALIGN.RIGHT, italic=True)


def role_box(slide, left, top, role, holder, width=BOX_W):
    """A chart box: the role name, and under it whoever holds it."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, BOX_H)
    shape.adjustments[0] = 0.12
    shape.fill.solid()
    shape.fill.fore_color.rgb = BOX_FILL
    shape.line.color.rgb = PURPLE
    shape.line.width = Pt(1.25)
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text(frame, role, 14, bold=True, colour=RGBColor(0x2B, 0x2B, 0x2B), align=PP_ALIGN.CENTER)
    if holder:
        p = frame.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = holder
        run.font.size = Pt(11)
        run.font.color.rgb = GREY
    return shape


def connect(slide, parent, child):
    """An elbow line from the bottom of `parent` to the top of `child`.

    The connector stays glued to both boxes, so PowerPoint reroutes it when
    either one moves. Its initial frame is written by hand: an elbow connector
    runs horizontal first, so a top-down tree needs it turned 90 degrees, with
    a vertical flip when the child sits to the right of its parent.
    """
    line = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, 0, 0)
    line.begin_connect(parent, BOTTOM)
    line.end_connect(child, TOP)
    line.line.color.rgb = LINE
    line.line.width = Pt(1.25)

    x1, y1 = parent.left + parent.width // 2, parent.top + parent.height
    x2, y2 = child.left + child.width // 2, child.top
    width, height = abs(y2 - y1), abs(x2 - x1)
    xfrm = line._element.spPr.get_or_add_xfrm()
    for attribute in ("flipH", "flipV"):
        xfrm.attrib.pop(attribute, None)
    xfrm.set("rot", "5400000")
    if x2 > x1:
        xfrm.set("flipV", "1")
    xfrm.off.x = (x1 + x2) // 2 - width // 2
    xfrm.off.y = (y1 + y2) // 2 - height // 2
    xfrm.ext.cx, xfrm.ext.cy = width, height
    return line


def row(count, width, gap):
    """Left positions of `count` boxes of `width`, centred on the slide."""
    total = count * width + (count - 1) * gap
    start = (SLIDE_W - total) // 2
    return [start + index * (width + gap) for index in range(count)]


def chart(slide, top, branches, leaves, top_y=Cm(4.6)):
    """One top box, three branches, two roles under each. Returns the branch shapes."""
    level_gap = Cm(2.2)
    root = role_box(slide, row(1, BOX_W, 0)[0], top_y, *top)

    branch_y = top_y + BOX_H + level_gap
    branch_shapes = []
    for left, branch in zip(row(3, BOX_W, Cm(5.9)), branches):
        shape = role_box(slide, left, branch_y, *branch)
        connect(slide, root, shape)
        branch_shapes.append(shape)

    leaf_y = branch_y + BOX_H + level_gap
    for index, left in enumerate(row(6, BOX_W, Cm(0.6))):
        shape = role_box(slide, left, leaf_y, *leaves[index])
        connect(slide, branch_shapes[index // 2], shape)
    return branch_shapes


def link_button(slide, label, target):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, SLIDE_H - Cm(2.3), Cm(6.5), Cm(1.1)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PURPLE
    shape.line.fill.background()
    shape.shadow.inherit = False
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text(shape.text_frame, label, 12, bold=True, colour=WHITE, align=PP_ALIGN.CENTER)
    shape.click_action.target_slide = target


def role_list(slide, columns, examples, rows=9):
    table = slide.shapes.add_table(
        1 + len(examples) + rows, len(columns),
        MARGIN, Cm(4.4), SLIDE_W - 2 * MARGIN, Cm(12.5),
    ).table
    for index, header in enumerate(columns):
        cell = table.cell(0, index)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PURPLE
        cell.text_frame.text = ""
        text(cell.text_frame, header, 13, bold=True, colour=WHITE)
    for line in range(1, len(examples) + rows + 1):
        for index in range(len(columns)):
            cell = table.cell(line, index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if line % 2 else BOX_FILL
            cell.text_frame.text = ""
            value = examples[line - 1][index] if line <= len(examples) else ""
            text(cell.text_frame, value, 11, colour=GREY, italic=True)


def build(t):
    deck = Presentation()
    deck.slide_width, deck.slide_height = SLIDE_W, SLIDE_H
    blank = deck.slide_layouts[6]
    slides = [deck.slides.add_slide(blank) for _ in range(6)]

    title(slides[0], t["list_heading"], t["list_hint"])
    role_list(slides[0], t["list_cols"], t["list_examples"])
    footer(slides[0], t["footer"])

    example = t["example"]
    title(slides[1], t["example_heading"], t["example_hint"])
    branch_shapes = chart(slides[1], example["top"], example["branches"], example["leaves"])
    for shape, target in zip(branch_shapes, slides[2:5]):
        shape.click_action.target_slide = target
    footer(slides[1], t["footer"])

    for index, slide in enumerate(slides[2:5]):
        branch = example["branches"][index]
        title(slide, branch[0], t["branch_hint"])
        parent = role_box(slide, row(1, BOX_W, 0)[0], Cm(5.2), *branch)
        for left, leaf in zip(row(2, Cm(6.5), Cm(3)), example["leaves"][index * 2:index * 2 + 2]):
            child = role_box(slide, left, Cm(10), *leaf, width=Cm(6.5))
            connect(slide, parent, child)
        link_button(slide, t["back"], slides[1])
        footer(slide, t["footer"])

    b = t["blank"]
    title(slides[5], t["blank_heading"], t["blank_hint"])
    chart(slides[5], (b["role"], b["holder"]), [(b["role"], b["holder"])] * 3,
          [(b["role"], b["holder"])] * 6)
    footer(slides[5], t["footer"])
    return deck


TEMPLATES = {
    "en": ("org-chart-template.pptx", {
        "list_heading": "Step 1. List the roles",
        "list_hint": "One line per role, before drawing anything. A person can hold several "
                     "roles, so write the roles first and the names after.",
        "list_cols": ["Role", "Reports to", "Held by", "Decides alone on"],
        "list_examples": [
            ["Customer Support", "Operations", "Ana", "Refunds under 200 euros"],
            ["Partnerships", "Sales", "vacant", "Reseller contracts"],
        ],
        "example_heading": "Step 2. The chart, filled in",
        "example_hint": "Ten roles, eight people. Type your roles over the boxes. In slide "
                        "show, click a branch to open its slide.",
        "example": {
            "top": ("Managing Director", "Sarah"),
            "branches": [
                ("Operations", "Malik"),
                ("Sales", "Ines"),
                ("Product", "Tom"),
            ],
            "leaves": [
                ("Customer Support", "Ana"),
                ("Finance and Admin", "Malik"),
                ("Account Management", "Ines, Jo"),
                ("Partnerships", "vacant"),
                ("Engineering", "Tom, Lea"),
                ("Design", "Noah"),
            ],
        },
        "branch_hint": "One slide per branch keeps the text readable from the back of the "
                       "room. Add the roles of this team here.",
        "back": "Back to the org chart",
        "blank_heading": "Step 3. The blank chart",
        "blank_hint": "Move a box and its lines follow. Copy a box with Ctrl+D, then drag "
                      "the ends of a connector onto it.",
        "blank": {"role": "Role", "holder": "Name"},
        "footer": "Template by Rolebase, rolebase.io",
    }),
    "fr": ("modele-organigramme.pptx", {
        "list_heading": "Étape 1. Listez les rôles",
        "list_hint": "Une ligne par rôle, avant de dessiner quoi que ce soit. Une personne "
                     "peut tenir plusieurs rôles, donc écrivez les rôles d’abord et les noms "
                     "ensuite.",
        "list_cols": ["Rôle", "Rattaché à", "Tenu par", "Décide seul de"],
        "list_examples": [
            ["Support client", "Opérations", "Ana", "Les remboursements sous 200 euros"],
            ["Partenariats", "Commercial", "vacant", "Les contrats revendeurs"],
        ],
        "example_heading": "Étape 2. L’organigramme rempli",
        "example_hint": "Dix rôles, huit personnes. Écrivez vos rôles par-dessus les cases. "
                        "En mode diaporama, cliquez sur une branche pour ouvrir sa diapositive.",
        "example": {
            "top": ("Direction générale", "Sarah"),
            "branches": [
                ("Opérations", "Malik"),
                ("Commercial", "Inès"),
                ("Produit", "Tom"),
            ],
            "leaves": [
                ("Support client", "Ana"),
                ("Finance et admin", "Malik"),
                ("Gestion de comptes", "Inès, Jo"),
                ("Partenariats", "vacant"),
                ("Développement", "Tom, Léa"),
                ("Design", "Noah"),
            ],
        },
        "branch_hint": "Une diapositive par branche garde le texte lisible depuis le fond "
                       "de la salle. Ajoutez ici les rôles de cette équipe.",
        "back": "Retour à l’organigramme",
        "blank_heading": "Étape 3. L’organigramme vierge",
        "blank_hint": "Déplacez une case et ses traits suivent. Dupliquez une case avec "
                      "Ctrl+D, puis faites glisser les extrémités d’un connecteur dessus.",
        "blank": {"role": "Rôle", "holder": "Prénom"},
        "footer": "Modèle proposé par Rolebase, rolebase.io",
    }),
}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for lang, (filename, t) in TEMPLATES.items():
        path = OUT / filename
        build(t).save(path)
        print(f"{lang}: {path}")


if __name__ == "__main__":
    main()
