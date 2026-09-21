#!/usr/bin/env python3
"""Generate the Word org chart templates offered on the "org chart in Word" post.

Lives beside the page it serves: `en.mdx` and `fr.mdx` in this folder,
published on /en/blog/org-chart-word and /fr/blog/org-chart-word.

Files produced, under `public/downloads/` and served on `/downloads/<file>`
    org-chart-template.docx   EN, linked from en.mdx
    modele-organigramme.docx  FR, linked from fr.mdx

Each document carries a portrait page to list the roles before drawing
anything, then two landscape pages holding a three-level chart: a filled
example and a blank one to type over. The chart is laid out as a Word table,
with the boxes and the connecting lines drawn as cell borders, so the reader
edits it by typing rather than by dragging shapes.

These binaries are build artefacts committed to the repo. Edit the definitions
below and re-run this script rather than editing the files themselves.

    pip install python-docx
    python3 src/content/blog/org-chart-word/generate-org-chart-template.py
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import (
    WD_ALIGN_VERTICAL,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[4] / "public" / "downloads"

PURPLE = RGBColor(0x98, 0x70, 0xF0)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
BOX_FILL = "F3EDFE"
BOX_LINE = "9870F0"
LINK_LINE = "9A9A9A"

COLS = 12          # the chart grid: 3 branches of 4 columns
PREP_ROWS = 12     # blank lines on the role list
BOX_HEIGHT = Cm(1.5)
LINK_HEIGHT = Cm(0.35)


def styled(doc):
    """Base typography shared by every template."""
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name, size in (("Heading 1", 20), ("Heading 2", 14)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = PURPLE
        st.font.bold = True
    return doc


def hint(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = GREY
    return p


def landscape(doc):
    """Open a landscape section, wide enough for a three-level chart."""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    width, height = section.page_width, section.page_height
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = height, width
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Cm(1.5))
    return section


def borders(cell, **sides):
    """Draw only the named sides. `sides` maps top/left/bottom/right to a colour."""
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    element = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        colour = sides.get(side)
        edge.set(qn("w:val"), "single" if colour else "nil")
        if colour:
            edge.set(qn("w:sz"), "8")
            edge.set(qn("w:color"), colour)
        element.append(edge)
    tc_pr.append(element)


def shade(cell, colour):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), colour)
    tc_pr.append(element)


def box(cell, title, holder=""):
    """A chart box: the role name, and under it whoever holds it."""
    borders(cell, top=BOX_LINE, left=BOX_LINE, bottom=BOX_LINE, right=BOX_LINE)
    shade(cell, BOX_FILL)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(11)
    if holder:
        second = cell.add_paragraph()
        second.alignment = WD_ALIGN_PARAGRAPH.CENTER
        second.paragraph_format.space_after = Pt(0)
        run = second.add_run(holder)
        run.font.size = Pt(9)
        run.font.color.rgb = GREY


def fixed(row, height):
    row.height = height
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY


def link_row(row, tops=(), rights=()):
    """One row of connecting lines, drawn as cell borders."""
    for column in set(tops) | set(rights):
        sides = {}
        if column in tops:
            sides["top"] = LINK_LINE
        if column in rights:
            sides["right"] = LINK_LINE
        borders(row.cells[column], **sides)


def chart(doc, top, branches, leaves):
    """A three-level chart: one top box, three branches, two roles under each.

    The grid holds 12 columns. Boxes span four columns on the branch row and
    two on the leaf row, which puts every box centre on a column boundary,
    where a cell border can draw the line that reaches it.
    """
    table = doc.add_table(rows=7, cols=COLS)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            borders(cell)
            cell.width = Cm(2.05)

    fixed(table.rows[0], BOX_HEIGHT)
    box(table.rows[0].cells[4].merge(table.rows[0].cells[7]), *top)

    fixed(table.rows[1], LINK_HEIGHT)
    link_row(table.rows[1], rights=(5,))

    fixed(table.rows[2], LINK_HEIGHT)
    link_row(table.rows[2], tops=range(2, 10), rights=(1, 5, 9))

    fixed(table.rows[3], BOX_HEIGHT)
    for index, branch in enumerate(branches):
        start = index * 4
        box(table.rows[3].cells[start].merge(table.rows[3].cells[start + 3]), *branch)

    fixed(table.rows[4], LINK_HEIGHT)
    link_row(table.rows[4], rights=(1, 5, 9))

    fixed(table.rows[5], LINK_HEIGHT)
    centres = (1, 5, 9)
    link_row(
        table.rows[5],
        tops=[column for centre in centres for column in (centre, centre + 1)],
        rights=[column for centre in centres for column in (centre - 1, centre + 1)],
    )

    fixed(table.rows[6], BOX_HEIGHT)
    for index, leaf in enumerate(leaves):
        start = index * 2
        box(table.rows[6].cells[start].merge(table.rows[6].cells[start + 1]), *leaf)
    return table


def prep_table(doc, columns, examples):
    table = doc.add_table(rows=1 + len(examples) + PREP_ROWS, cols=len(columns))
    table.style = "Table Grid"
    for index, header in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = ""
        cell.paragraphs[0].add_run(header).bold = True
    for line, example in enumerate(examples, start=1):
        for index, value in enumerate(example):
            cell = table.rows[line].cells[index]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.italic = True
            run.font.color.rgb = GREY
    return table


def footer_note(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def build(t):
    doc = styled(Document())
    doc.add_heading(t["title"], level=1)
    hint(doc, t["intro"])

    doc.add_heading(t["list_heading"], level=2)
    hint(doc, t["list_hint"])
    prep_table(doc, t["list_cols"], t["list_examples"])

    landscape(doc)
    doc.add_heading(t["example_heading"], level=2)
    hint(doc, t["example_hint"])
    chart(doc, t["example"]["top"], t["example"]["branches"], t["example"]["leaves"])

    doc.add_page_break()
    doc.add_heading(t["blank_heading"], level=2)
    hint(doc, t["blank_hint"])
    blank_top = (t["blank"]["top"], t["blank"]["holder"])
    chart(doc,
          blank_top,
          [(t["blank"]["branch"], t["blank"]["holder"])] * 3,
          [(t["blank"]["leaf"], t["blank"]["holder"])] * 6)
    footer_note(doc, t["footer"])
    return doc


EN_FOOTER = "Template by Rolebase, rolebase.io"
FR_FOOTER = "Modèle proposé par Rolebase, rolebase.io"

TEMPLATES = {
    "en": ("org-chart-template.docx", {
        "title": "Org chart template",
        "intro": "List your roles on the first page, then type them into the chart on the "
                 "landscape pages. Each box is a table cell, so you type over it. To add a "
                 "branch, copy a column. The whole chart prints on one page.",
        "list_heading": "Step 1. List the roles",
        "list_hint": "Fill in one line per role, before drawing anything. A person can hold "
                     "several roles, so write the roles first and the names after.",
        "list_cols": ["Role", "Reports to", "Held by", "Decides alone on"],
        "list_examples": [
            ["Customer Support", "Operations", "Ana", "Refunds under 200 euros"],
            ["Partnerships", "Sales", "vacant", "Reseller contracts"],
        ],
        "example_heading": "Step 2. The chart, filled in",
        "example_hint": "This example has ten roles and eight people. Type your own roles "
                        "over the boxes. Delete the extra boxes.",
        "example": {
            "top": ("Managing Director", "Sarah"),
            "branches": [
                ("Operations", "Malik"),
                ("Sales", "Ines"),
                ("Product", "Tom"),
            ],
            "leaves": [
                ("Customer Support", "Ana"),
                ("Finance and Admin", "Malik"),
                ("Account Management", "Ines, Jo"),
                ("Partnerships", "vacant"),
                ("Engineering", "Tom, Lea"),
                ("Design", "Noah"),
            ],
        },
        "blank_heading": "Step 3. The blank chart",
        "blank_hint": "This chart repeats the same layout, ready to fill in. Select "
                      "a column and copy it to add a branch. Delete a row of boxes to "
                      "drop a level.",
        "blank": {
            "top": "Role",
            "branch": "Role",
            "leaf": "Role",
            "holder": "Name",
        },
        "footer": EN_FOOTER,
    }),
    "fr": ("modele-organigramme.docx", {
        "title": "Modèle d’organigramme",
        "intro": "Listez vos rôles sur la première page, puis reportez-les dans l’organigramme "
                 "des pages en paysage. Chaque case est une cellule de tableau, donc vous "
                 "écrivez par-dessus. Pour ajouter une branche, copiez une colonne. "
                 "L’organigramme entier tient sur une page à l’impression.",
        "list_heading": "Étape 1. Listez les rôles",
        "list_hint": "Remplissez une ligne par rôle, avant de dessiner quoi que ce soit. Une "
                     "personne peut tenir plusieurs rôles, donc écrivez les rôles d’abord et les "
                     "noms ensuite.",
        "list_cols": ["Rôle", "Rattaché à", "Tenu par", "Décide seul de"],
        "list_examples": [
            ["Support client", "Opérations", "Ana", "Les remboursements sous 200 euros"],
            ["Partenariats", "Commercial", "vacant", "Les contrats revendeurs"],
        ],
        "example_heading": "Étape 2. L’organigramme rempli",
        "example_hint": "Cet exemple compte dix rôles et huit personnes. Écrivez vos propres "
                        "rôles par-dessus les cases. Supprimez les cases en trop.",
        "example": {
            "top": ("Direction générale", "Sarah"),
            "branches": [
                ("Opérations", "Malik"),
                ("Commercial", "Inès"),
                ("Produit", "Tom"),
            ],
            "leaves": [
                ("Support client", "Ana"),
                ("Finance et admin", "Malik"),
                ("Gestion de comptes", "Inès, Jo"),
                ("Partenariats", "vacant"),
                ("Développement", "Tom, Léa"),
                ("Design", "Noah"),
            ],
        },
        "blank_heading": "Étape 3. L’organigramme vierge",
        "blank_hint": "Cet organigramme reprend la même disposition, prêt à remplir. "
                      "Sélectionnez une colonne et copiez-la pour ajouter une branche. "
                      "Supprimez une ligne de cases pour retirer un niveau.",
        "blank": {
            "top": "Rôle",
            "branch": "Rôle",
            "leaf": "Rôle",
            "holder": "Prénom",
        },
        "footer": FR_FOOTER,
    }),
}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for lang, (filename, t) in TEMPLATES.items():
        path = OUT / filename
        build(t).save(path)
        print(f"{lang}: {path}")


if __name__ == "__main__":
    main()
