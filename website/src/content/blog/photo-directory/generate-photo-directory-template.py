#!/usr/bin/env python3
"""Generate the Word photo directory templates offered on the "photo directory" post.

Lives beside the page it serves: `en.mdx` and `fr.mdx` in this folder,
published on /en/blog/photo-directory and /fr/blog/photo-directory.

Files produced, under `public/downloads/` and served on `/downloads/<file>`
    photo-directory-template.docx   EN, linked from en.mdx
    modele-trombinoscope.docx       FR, linked from fr.mdx

Each document carries three portrait pages: a filled example using the demo
avatars of the site (`public/demo-avatars/`), a blank grid of twelve cards to
fill in, and a photo consent form to hand to each person before their photo
goes in. The cards are laid out as a Word table, so the reader replaces a
photo by clicking in its cell and inserting a picture, and types over the rest.

These binaries are build artefacts committed to the repo. Edit the definitions
below and re-run this script rather than editing the files themselves.

    pip install python-docx
    python3 src/content/blog/photo-directory/generate-photo-directory-template.py
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[4]
OUT = ROOT / "public" / "downloads"
AVATARS = ROOT / "public" / "demo-avatars"

PURPLE = RGBColor(0x98, 0x70, 0xF0)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
CARD_LINE = "D9CCF9"
PHOTO_FILL = "F3EDFE"

COLUMNS = 3
CARD_WIDTH = Cm(5.6)
GAP_WIDTH = Cm(0.6)
PHOTO_HEIGHT = Cm(3.1)
TEXT_HEIGHT = Cm(1.7)
SPACER_HEIGHT = Cm(0.4)
PHOTO_SIZE = Cm(2.6)
BLANK_CARDS = 12

TEXTS = {
    "fr": {
        "file": "modele-trombinoscope.docx",
        "title": "Trombinoscope de [nom de l’entreprise]",
        "updated": "Mis à jour le [date]. Prochaine relecture le [date].",
        "example_hint": "Voici un exemple rempli. Remplacez chaque photo en cliquant dans sa case, puis Insertion > Images. Une personne qui refuse la photo garde sa carte, avec ses initiales à la place.",
        "blank_title": "Trombinoscope à remplir",
        "blank_hint": "Remplissez une carte par personne, avec son nom, ses rôles et son équipe. Copiez une ligne du tableau pour ajouter des cartes.",
        "photo": "Photo",
        "name": "Prénom Nom",
        "roles": "Rôles",
        "team": "Équipe",
        "people": [
            ("alice.jpg", "Alice", "Coordination produit", "Produit"),
            ("bruno.jpg", "Bruno", "Développement back-end, astreinte", "Tech"),
            ("chloe.jpg", "Chloé", "Design produit", "Produit"),
            ("emma.jpg", "Emma", "Ventes grands comptes", "Commercial"),
            ("tom.jpg", "Tom", "Support client, accueil des nouveaux", "Service client"),
            ("camille.jpg", "Camille", "Finance, paie", "Direction"),
        ],
        "consent_title": "Autorisation d’utiliser ma photo dans le trombinoscope",
        "consent_hint": "Modèle à adapter et à faire relire par la personne chargée des données personnelles dans votre entreprise.",
        "consent_body": [
            "Je soussigné(e) [prénom nom], [rôle ou poste], autorise [nom de l’entreprise] à utiliser ma photo dans le trombinoscope interne de l’entreprise.",
            "La photo sert uniquement à aider mes collègues à me reconnaître et à savoir qui fait quoi. Elle apparaît sur les supports suivants : [intranet, outil d’organigramme, document imprimé affiché dans les locaux]. Elle n’est utilisée pour aucun autre usage, en particulier aucune communication externe, sans un nouvel accord de ma part.",
            "La photo est retirée au plus tard [un mois] après mon départ de l’entreprise.",
            "Je peux retirer cette autorisation à tout moment, sans avoir à me justifier et sans conséquence pour moi, en écrivant à [adresse e-mail ou service]. Ma photo est alors retirée dans un délai de [quinze jours].",
        ],
        "consent_choice": [
            "☐ J’accepte que ma photo figure dans le trombinoscope interne.",
            "☐ Je refuse. Ma carte figure dans le trombinoscope avec mes initiales à la place de la photo.",
        ],
        "consent_sign": "Fait à [ville], le [date].        Signature :",
    },
    "en": {
        "file": "photo-directory-template.docx",
        "title": "[Company name] photo directory",
        "updated": "Updated on [date]. Next review on [date].",
        "example_hint": "Here is a filled example. Replace each photo by clicking in its cell, then Insert > Pictures. Anyone who declines a photo keeps their card, with their initials instead.",
        "blank_title": "Photo directory to fill in",
        "blank_hint": "Fill in one card per person, with their name, roles and team. Copy a table row to add more cards.",
        "photo": "Photo",
        "name": "First Last",
        "roles": "Roles",
        "team": "Team",
        "people": [
            ("alice.jpg", "Alice", "Product coordination", "Product"),
            ("bruno.jpg", "Bruno", "Back-end development, on-call", "Tech"),
            ("chloe.jpg", "Chloé", "Product design", "Product"),
            ("emma.jpg", "Emma", "Key account sales", "Sales"),
            ("tom.jpg", "Tom", "Customer support, new hire buddy", "Customer service"),
            ("camille.jpg", "Camille", "Finance, payroll", "Leadership"),
        ],
        "consent_title": "Permission to use my photo in the photo directory",
        "consent_hint": "A template to adapt and to have checked by whoever handles personal data in your company.",
        "consent_body": [
            "I, [first and last name], [role or position], allow [company name] to use my photo in the company’s internal photo directory.",
            "The photo is only used to help my colleagues recognise me and know who does what. It appears on the following media: [intranet, org chart tool, printed sheet displayed in the office]. It is not used for anything else, in particular any external communication, without my further agreement.",
            "The photo is removed no later than [one month] after I leave the company.",
            "I can withdraw this permission at any time, without giving a reason and with no consequence for me, by writing to [email address or department]. My photo is then removed within [two weeks].",
        ],
        "consent_choice": [
            "☐ I agree to my photo appearing in the internal photo directory.",
            "☐ I decline. My card appears in the directory with my initials instead of a photo.",
        ],
        "consent_sign": "Signed in [city], on [date].        Signature:",
    },
}


def styled(doc):
    """Base typography, and 1.5 cm margins so three cards fit across."""
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name, size in (("Heading 1", 20), ("Heading 2", 14)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = PURPLE
        st.font.bold = True
    section = doc.sections[0]
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Cm(1.5))
    return doc


def hint(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = GREY
    return p


def borders(cell, **sides):
    """Draw only the named sides. `sides` maps top/left/bottom/right to a colour."""
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    element = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        colour = sides.get(side)
        edge.set(qn("w:val"), "single" if colour else "nil")
        if colour:
            edge.set(qn("w:sz"), "8")
            edge.set(qn("w:color"), colour)
        element.append(edge)
    tc_pr.append(element)


def shade(cell, colour):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), colour)
    tc_pr.append(element)


def fixed(row, height):
    row.height = height
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY


def line(cell, text, first=False, bold=False, size=10, colour=None):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if colour:
        run.font.color.rgb = colour


def photo_cell(cell, picture, placeholder):
    borders(cell, top=CARD_LINE, left=CARD_LINE, right=CARD_LINE)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if picture:
        p.add_run().add_picture(str(picture), width=PHOTO_SIZE, height=PHOTO_SIZE)
    else:
        shade(cell, PHOTO_FILL)
        run = p.add_run(placeholder)
        run.font.color.rgb = GREY


def text_cell(cell, name, roles, team, muted):
    borders(cell, left=CARD_LINE, right=CARD_LINE, bottom=CARD_LINE)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    line(cell, name, first=True, bold=True, size=11, colour=GREY if muted else None)
    line(cell, roles, size=9, colour=GREY)
    line(cell, team, size=9, colour=PURPLE)


def cards(doc, entries, placeholder):
    """A grid of cards, three across. Each entry is (picture, name, roles, team)."""
    rows_of_cards = [entries[i : i + COLUMNS] for i in range(0, len(entries), COLUMNS)]
    table = doc.add_table(rows=0, cols=COLUMNS * 2 - 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Fixed layout, so Word keeps the card widths instead of fitting the text
    table.autofit = False
    widths = [CARD_WIDTH if i % 2 == 0 else GAP_WIDTH for i in range(COLUMNS * 2 - 1)]
    for column, width in zip(table.columns, widths):
        column.width = width
    # An explicit total width, otherwise Word still spreads the columns
    # according to their content
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(w.twips for w in widths)))

    for index, group in enumerate(rows_of_cards):
        photo_row, text_row = table.add_row(), table.add_row()
        fixed(photo_row, PHOTO_HEIGHT)
        fixed(text_row, TEXT_HEIGHT)
        rows = [photo_row, text_row]
        if index < len(rows_of_cards) - 1:
            spacer = table.add_row()
            fixed(spacer, SPACER_HEIGHT)
            rows.append(spacer)
        for row in rows:
            for cell, width in zip(row.cells, widths):
                cell.width = width
                borders(cell)
        for position, (picture, name, roles, team) in enumerate(group):
            column = position * 2
            photo_cell(photo_row.cells[column], picture, placeholder)
            text_cell(text_row.cells[column], name, roles, team, muted=picture is None)
    return table


def build(lang):
    t = TEXTS[lang]
    doc = styled(Document())

    doc.add_heading(t["title"], level=1)
    hint(doc, t["updated"])
    example = [(AVATARS / f, name, roles, team) for f, name, roles, team in t["people"]]
    cards(doc, example, t["photo"])
    doc.add_paragraph()
    hint(doc, t["example_hint"])

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading(t["blank_title"], level=1)
    hint(doc, t["blank_hint"])
    blank = [(None, t["name"], t["roles"], t["team"])] * BLANK_CARDS
    cards(doc, blank, t["photo"])

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading(t["consent_title"], level=1)
    hint(doc, t["consent_hint"])
    for paragraph in t["consent_body"]:
        doc.add_paragraph(paragraph).paragraph_format.space_after = Pt(10)
    for choice in t["consent_choice"]:
        doc.add_paragraph(choice).paragraph_format.space_after = Pt(6)
    doc.add_paragraph()
    doc.add_paragraph(t["consent_sign"])

    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / t["file"]
    doc.save(path)
    print(f"wrote {path.relative_to(ROOT)}")


if __name__ == "__main__":
    for lang in TEXTS:
        build(lang)
